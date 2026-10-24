"""Turn the core attraction review sheet into config/core_attractions.json.

    .\\.venv\\Scripts\\python.exe scripts\\import_core_review.py

Then rebuild the catalogue so the app uses it:

    .\\.venv\\Scripts\\python.exe scripts\\add_overture_to_catalogue.py

Only rows whose status is not "待核验" are exported. Entrances without coordinates are
skipped. The JSON is the reviewed source of truth that goes into Git; the xlsx is the
working copy people edit.
"""
from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
BOOK = ROOT / 'review' / 'core_attractions_review.xlsx'
OUT = ROOT / 'config' / 'core_attractions.json'

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit('Please run: .\\.venv\\Scripts\\python.exe -m pip install openpyxl')

# Decisions made while reviewing the data (see review notes). Kept here so they are versioned.
CORRECTIONS = {
    # NHB point for Fort Siloso is on the mainland, about 3.5 km from the real site on Sentosa.
    'place:nhb_monuments:1684': {'coordinate_from_entrance': 'Fort Siloso',
                                 'reason': 'NHB 坐标在本岛，离圣淘沙实际位置约 3.5 km；改用 OneMap 位置'},
}
TITLE_OVERRIDES = {
    'place:stb_attractions:987': 'Singapore Oceanarium',
}
HIDE = {
    # Overture still has the old Jurong Bird Park site; the park moved to Mandai as Bird Paradise.
    'place:overture_places:d50564ad-bee0-4ea7-acaa-38b9770c2a21': 'Jurong Bird Park 已搬迁到万礼，改名 Bird Paradise',
    # STB record of the old name; River Wonders (Overture) is the current record.
    'place:stb_attractions:990': 'River Safari 已改名 River Wonders，使用新记录',
    # Same name as the theme park but placed in the city centre (~4.5 km away); likely a sales office or bad point.
    'place:overture_places:7c6d11c8-244a-4526-aade-b2afe9dc1231': '与环球影城同名但坐标在市中心，疑为售票点或坐标错误',
}
BASIS = {'官方': 'official', '估计': 'estimate', '人工核验': 'field_checked'}


def text(v):
    v = '' if v is None else str(v).strip()
    return v or None


def minutes(v, label, name):
    if v in (None, ''):
        return None
    try:
        n = int(float(v))
    except (TypeError, ValueError):
        sys.exit(f'{name}: {label} must be a number of minutes, got {v!r}')
    if not 0 < n <= 720:
        sys.exit(f'{name}: {label} must be between 1 and 720 minutes, got {n}')
    return n


def main():
    if not BOOK.is_file():
        sys.exit(f'Not found: {BOOK}')
    wb = load_workbook(BOOK, data_only=True)
    ws, es = wb['核心景点核验'], wb['入口']
    h = [c.value for c in ws[1]]
    col = {n: i for i, n in enumerate(h)}
    eh = [c.value for c in es[1]]
    ecol = {n: i for i, n in enumerate(eh)}

    entrances = {}
    for r in es.iter_rows(min_row=2, values_only=True):
        eid, name, lat, lon = r[ecol['地点 ID']], r[ecol['入口名称']], r[ecol['纬度']], r[ecol['经度']]
        if not eid or not name or lat in (None, '') or lon in (None, ''):
            continue
        lat, lon = float(lat), float(lon)
        if not (1.15 <= lat <= 1.48 and 103.6 <= lon <= 104.1):
            sys.exit(f'Entrance outside Singapore: {name} ({lat}, {lon})')
        entrances.setdefault(eid, []).append({
            'name': text(name), 'latitude': round(lat, 7), 'longitude': round(lon, 7),
            'type': text(r[ecol['入口类型']]), 'source': text(r[ecol['来源']]),
            'status': text(r[ecol['核验状态']]), 'note': text(r[ecol['备注']]),
        })

    places, skipped = {}, []
    for r in ws.iter_rows(min_row=2, values_only=True):
        eid, name = r[col['地点 ID']], r[col['景点']]
        if not eid:
            continue
        status = text(r[col['核验状态']])
        if status in (None, '待核验'):
            skipped.append(name)
            continue
        short, regular, long_ = (minutes(r[col['时长·短（分钟）']], '时长·短', name),
                                 minutes(r[col['时长·一般（分钟）']], '时长·一般', name),
                                 minutes(r[col['时长·长（分钟）']], '时长·长', name))
        duration = None
        if short and regular and long_:
            if not short <= regular <= long_:
                sys.exit(f'{name}: durations must be short <= regular <= long')
            duration = {'short': short, 'regular': regular, 'long': long_,
                        'basis': BASIS.get(text(r[col['时长依据']]), 'estimate')}
        places[eid] = {
            'name': name, 'priority': text(r[col['优先级']]),
            'opening_hours': text(r[col['开放时间']]), 'last_admission': text(r[col['最后入场']]),
            'closed': text(r[col['休息日 / 闭馆']]), 'price': text(r[col['收费']]),
            'suggested_duration': duration, 'duration_note': text(r[col['时长说明']]),
            'entrances': entrances.get(eid, []),
            'source': text(r[col['来源链接']]), 'note': text(r[col['备注 / 问题']]),
            'status': status, 'checked_on': text(r[col['核验日期']]), 'checked_by': text(r[col['核验人']]),
        }
        if eid in CORRECTIONS:
            places[eid]['correction'] = CORRECTIONS[eid]
        if eid in TITLE_OVERRIDES:
            places[eid]['display_title'] = TITLE_OVERRIDES[eid]

    out = {'schema_version': 1, 'exported_on': date.today().isoformat(),
           'source_workbook': 'review/core_attractions_review.xlsx',
           'note': '资料核对（Claude）= 按官网或公开资料整理，未实地核验；时长 basis=estimate 为估计值。',
           'places': places, 'hide': HIDE}
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding='utf-8')
    n_ent = sum(len(p['entrances']) for p in places.values())
    n_dur = sum(1 for p in places.values() if p['suggested_duration'])
    n_hours = sum(1 for p in places.values() if p['opening_hours'])
    print(f'Exported {len(places)} places ({n_hours} with opening hours, {n_dur} with durations, {n_ent} entrances)')
    if skipped:
        print(f'Skipped {len(skipped)} rows still marked 待核验: {", ".join(map(str, skipped))}')
    print(f'Saved {OUT.relative_to(ROOT)}. Next: python scripts/add_overture_to_catalogue.py')


if __name__ == '__main__':
    main()
